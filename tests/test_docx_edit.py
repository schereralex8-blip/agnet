"""Typing answers into a Word document in place."""

import pytest

from homework_agent.tools.base import ToolError
from homework_agent.tools.docx_edit import FillDocumentTool, replace_in_paragraph

docx_lib = pytest.importorskip("docx", reason="python-docx not installed")


@pytest.fixture
def workspace(tmp_path):
    return tmp_path


def build(path, blocks):
    """blocks: str -> paragraph, (style, text) -> styled paragraph, [runs] -> split runs."""
    document = docx_lib.Document()
    for block in blocks:
        if isinstance(block, tuple):
            document.add_paragraph(block[1], style=block[0])
        elif isinstance(block, list):
            paragraph = document.add_paragraph()
            for run in block:
                paragraph.add_run(run)
        else:
            document.add_paragraph(block)
    document.save(str(path))
    return path


def paragraphs(path):
    return [p.text for p in docx_lib.Document(str(path)).paragraphs if p.text.strip()]


def styles(path):
    return [(p.style.name, p.text) for p in docx_lib.Document(str(path)).paragraphs if p.text.strip()]


def test_answer_lands_directly_under_its_question(workspace):
    path = build(workspace / "pset.docx", ["Q1: Differentiate x^3.", "Q2: Integrate 2x."])
    FillDocumentTool(workspace).run(
        path="pset.docx", edits=[{"anchor": "Q1: Differentiate x^3.", "answer": "3x^2"}]
    )
    assert paragraphs(path) == ["Q1: Differentiate x^3.", "3x^2", "Q2: Integrate 2x."]


def test_multi_line_answers_become_separate_paragraphs(workspace):
    path = build(workspace / "pset.docx", ["Q1: Solve it."])
    FillDocumentTool(workspace).run(
        path="pset.docx",
        edits=[{"anchor": "Q1: Solve it.", "answer": "Step 1: factor.\nStep 2: solve.\n\nx = 3"}],
    )
    assert paragraphs(path) == ["Q1: Solve it.", "Step 1: factor.", "Step 2: solve.", "x = 3"]


def test_a_blank_split_across_runs_is_filled(workspace):
    """Word stores a visible '________' as several runs - the match must span them."""
    path = build(
        workspace / "pset.docx",
        [["Q1: The derivative of sin(x) is ", "____", "____", "."]],
    )
    FillDocumentTool(workspace).run(
        path="pset.docx",
        edits=[{"anchor": "________", "answer": "cos(x)", "where": "replace"}],
    )
    assert paragraphs(path) == ["Q1: The derivative of sin(x) is cos(x)."]


def test_replacing_keeps_the_rest_of_the_paragraph_formatting(workspace):
    path = workspace / "pset.docx"
    document = docx_lib.Document()
    paragraph = document.add_paragraph()
    bold = paragraph.add_run("Q1 (bold label): ")
    bold.bold = True
    paragraph.add_run("Answer:")
    document.save(str(path))

    FillDocumentTool(workspace).run(
        path="pset.docx", edits=[{"anchor": "Answer:", "answer": "42", "where": "replace"}]
    )
    reloaded = docx_lib.Document(str(path)).paragraphs[0]
    assert reloaded.text == "Q1 (bold label): 42"
    assert reloaded.runs[0].bold is True


def test_append_continues_the_question_paragraph(workspace):
    path = build(workspace / "pset.docx", ["Q1: The capital of France is"])
    FillDocumentTool(workspace).run(
        path="pset.docx",
        edits=[{"anchor": "capital of France", "answer": "Paris.", "where": "append"}],
    )
    assert paragraphs(path) == ["Q1: The capital of France is Paris."]


def test_answers_do_not_join_the_assignments_numbering(workspace):
    """An answer under a List Number question must not become the next question."""
    path = build(
        workspace / "pset.docx",
        [("List Number", "Q1: First question."), ("List Number", "Q2: Second question.")],
    )
    FillDocumentTool(workspace).run(
        path="pset.docx", edits=[{"anchor": "Q1: First question.", "answer": "Answer one."}]
    )
    assert styles(path) == [
        ("List Number", "Q1: First question."),
        ("Normal", "Answer one."),
        ("List Number", "Q2: Second question."),
    ]


def test_an_answer_under_a_heading_is_not_a_heading(workspace):
    path = workspace / "pset.docx"
    document = docx_lib.Document()
    document.add_heading("Q1: State the theorem.", 2)
    document.save(str(path))
    FillDocumentTool(workspace).run(
        path="pset.docx", edits=[{"anchor": "State the theorem", "answer": "Every f is bounded."}]
    )
    assert styles(path)[1] == ("Normal", "Every f is bounded.")


def test_indentation_is_carried_over(workspace):
    from docx.shared import Inches

    path = workspace / "pset.docx"
    document = docx_lib.Document()
    paragraph = document.add_paragraph("Q1 (a): sub-question.")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    document.save(str(path))

    FillDocumentTool(workspace).run(
        path="pset.docx", edits=[{"anchor": "sub-question", "answer": "Answer."}]
    )
    inserted = docx_lib.Document(str(path)).paragraphs[1]
    assert inserted.paragraph_format.left_indent == Inches(0.5)


def test_tables_can_be_filled(workspace):
    path = workspace / "pset.docx"
    document = docx_lib.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Question"
    table.cell(0, 1).text = "Your answer"
    table.cell(1, 0).text = "Q1: 7 x 8"
    table.cell(1, 1).text = "Answer:"
    document.save(str(path))

    FillDocumentTool(workspace).run(
        path="pset.docx", edits=[{"anchor": "Answer:", "answer": "56", "where": "replace"}]
    )
    assert docx_lib.Document(str(path)).tables[0].cell(1, 1).text == "56"


def test_the_original_is_backed_up_once(workspace):
    path = build(workspace / "pset.docx", ["Q1: a", "Q2: b"])
    tool = FillDocumentTool(workspace)
    tool.run(path="pset.docx", edits=[{"anchor": "Q1: a", "answer": "first"}])

    backup = workspace / "pset.original.docx"
    assert backup.exists()
    assert paragraphs(backup) == ["Q1: a", "Q2: b"]

    # A second pass must not overwrite the pristine copy with the filled version.
    tool.run(path="pset.docx", edits=[{"anchor": "Q2: b", "answer": "second"}])
    assert paragraphs(backup) == ["Q1: a", "Q2: b"]


def test_an_ambiguous_anchor_is_refused_before_anything_is_written(workspace):
    path = build(workspace / "pset.docx", ["Answer:", "Answer:"])
    with pytest.raises(ToolError, match="ambiguous"):
        FillDocumentTool(workspace).run(
            path="pset.docx", edits=[{"anchor": "Answer:", "answer": "x", "where": "replace"}]
        )
    assert paragraphs(path) == ["Answer:", "Answer:"]


def test_a_missing_anchor_shows_what_the_document_says(workspace):
    build(workspace / "pset.docx", ["Q1: Differentiate x^3."])
    with pytest.raises(ToolError, match="Q1: Differentiate"):
        FillDocumentTool(workspace).run(
            path="pset.docx", edits=[{"anchor": "Question 7", "answer": "x"}]
        )


def test_whitespace_differences_in_the_anchor_are_tolerated(workspace):
    path = build(workspace / "pset.docx", ["Q1:   Differentiate    x^3."])
    FillDocumentTool(workspace).run(
        path="pset.docx", edits=[{"anchor": "Q1: Differentiate x^3.", "answer": "3x^2"}]
    )
    assert "3x^2" in paragraphs(path)


def test_non_word_files_are_refused(workspace):
    (workspace / "answers.md").write_text("# answers")
    with pytest.raises(ToolError, match="not a Word document"):
        FillDocumentTool(workspace).run(
            path="answers.md", edits=[{"anchor": "answers", "answer": "x"}]
        )


def test_empty_edits_and_empty_answers_are_refused(workspace):
    build(workspace / "pset.docx", ["Q1: a"])
    tool = FillDocumentTool(workspace)
    with pytest.raises(ToolError, match="edits is empty"):
        tool.run(path="pset.docx", edits=[])
    with pytest.raises(ToolError, match="empty answer"):
        tool.run(path="pset.docx", edits=[{"anchor": "Q1: a", "answer": "   "}])


def test_a_bad_where_value_is_reported(workspace):
    build(workspace / "pset.docx", ["Q1: a"])
    with pytest.raises(ToolError, match="where must be"):
        FillDocumentTool(workspace).run(
            path="pset.docx", edits=[{"anchor": "Q1: a", "answer": "x", "where": "sideways"}]
        )


def test_editing_needs_approval():
    assert FillDocumentTool(".").requires_approval is True


def test_paths_cannot_escape_the_workspace(workspace):
    with pytest.raises(ToolError, match="outside the homework folder"):
        FillDocumentTool(workspace).run(
            path="../elsewhere.docx", edits=[{"anchor": "a", "answer": "b"}]
        )


def test_missing_python_docx_explains_the_install(workspace, monkeypatch):
    build(workspace / "pset.docx", ["Q1: a"])
    monkeypatch.setitem(__import__("sys").modules, "docx", None)
    with pytest.raises(ToolError, match=r"homework-agent\[docx\]"):
        FillDocumentTool(workspace).run(path="pset.docx", edits=[{"anchor": "Q1", "answer": "x"}])


def test_replace_in_paragraph_reports_a_miss(workspace):
    document = docx_lib.Document()
    paragraph = document.add_paragraph("nothing to see")
    assert replace_in_paragraph(paragraph, "absent", "x") is False


def test_describe_call_is_readable():
    described = FillDocumentTool(".").describe_call({"path": "hw.docx", "edits": [{}, {}]})
    assert described == "fill_document(hw.docx, 2 answer(s))"
