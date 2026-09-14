"""Workspace-confined file tools, so the agent can read the actual assignment."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator

from homework_agent.tools.base import Tool, ToolError

# Reading a 200MB file into the context window helps nobody.
MAX_READ_BYTES = 400_000
MAX_CHARS = 60_000
SKIP_DIRS = {".git", "__pycache__", "node_modules", ".venv", "venv", ".idea", ".DS_Store"}
TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".tex", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".py", ".ipynb", ".java", ".c", ".h", ".cpp", ".cs", ".js", ".ts", ".html", ".css",
    ".sql", ".r", ".m", ".sh", ".log", "",
}


def resolve_in_workspace(workspace: Path, relative: str) -> Path:
    """Resolve `relative` inside the workspace, refusing anything that escapes it."""
    workspace = workspace.resolve()
    candidate = (workspace / relative).expanduser()
    try:
        resolved = candidate.resolve()
    except OSError as exc:
        raise ToolError(f"cannot resolve path {relative!r}: {exc}") from None
    if resolved != workspace and workspace not in resolved.parents:
        raise ToolError(
            f"{relative!r} is outside the homework folder ({workspace}). "
            "Only paths inside it can be read or written."
        )
    return resolved


# Formats a student is likely to have, that are not plain text and are not PDF.
BINARY_FORMATS = {
    ".doc": "a Word document in the old binary format",
    ".odt": "an OpenDocument file",
    ".pptx": "a PowerPoint deck",
    ".xlsx": "an Excel workbook",
    ".zip": "a zip archive",
    ".png": "an image",
    ".jpg": "an image",
    ".jpeg": "an image",
    ".heic": "an image",
}


def _read_text(path: Path, shown: str) -> str:
    """Read a file as UTF-8 text, refusing binary rather than returning mojibake."""
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise ToolError(f"could not read {shown!r}: {exc}") from None

    described = BINARY_FORMATS.get(path.suffix.lower())
    if described:
        raise ToolError(
            f"{shown} is {described}, which cannot be read directly. "
            "Export it to PDF or plain text and point me at that, or paste the question in."
        )

    # NUL bytes decode happily as UTF-8 control characters, so check for them first -
    # they are the clearest sign the file is not text at all.
    if b"\x00" in raw[:8192]:
        raise ToolError(
            f"{shown} is a binary file, not something I can read as text. "
            "Export it to PDF or plain text, or paste the question in."
        )
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        # Latin-1 covers most of the rest: an assignment saved out of Word, say.
        return raw.decode("latin-1", errors="replace")


def _docx_body(document: Any) -> Iterator[Any]:
    """Yield paragraphs and tables in the order they appear in the document.

    python-docx exposes `.paragraphs` and `.tables` as separate lists, which loses
    their interleaving - and in a problem set the table usually belongs to the
    question right above it.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _render_docx_table(table: Any) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        raise ToolError(
            f"{path.name} is a Word document and python-docx is not installed. "
            "Install it with 'pip install homework-agent[docx]', or export the file to PDF "
            "or plain text."
        ) from None

    try:
        document = docx.Document(str(path))
        parts = []
        for block in _docx_body(document):
            if hasattr(block, "rows"):  # a table
                rendered = _render_docx_table(block)
            else:
                rendered = block.text.strip()
            if rendered.strip():
                parts.append(rendered)
    except ToolError:
        raise
    except Exception as exc:  # noqa: BLE001 - a bad zip raises any number of things
        raise ToolError(f"could not read Word document {path.name}: {exc}") from None

    if not parts:
        raise ToolError(
            f"{path.name} has no readable text - the questions may be images. "
            "Ask the student to type out the question."
        )
    return "\n\n".join(parts)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ToolError(
            f"{path.name} is a PDF and pypdf is not installed. "
            "Install it with 'pip install homework-agent[pdf]', or ask the student to paste "
            "the relevant question as text."
        ) from None
    try:
        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - malformed PDFs raise all sorts
        raise ToolError(f"could not read PDF {path.name}: {exc}") from None
    text = "\n\n".join(f"--- page {i} ---\n{body}" for i, body in enumerate(pages, 1) if body)
    if not text.strip():
        raise ToolError(
            f"{path.name} has no extractable text - it is probably a scan. "
            "Ask the student to type out the question."
        )
    return text


class ReadAssignmentTool(Tool):
    name = "read_assignment"
    description = (
        "Read a file from the student's homework folder: the problem set, their draft, their "
        "code, a data file. Always read the file before answering questions about 'the "
        "assignment' or a numbered question. Handles text files, PDFs, and Word documents."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Path relative to the homework folder, e.g. 'pset3.pdf'.",
            }
        },
        "required": ["path"],
    }

    def __init__(self, workspace: Path) -> None:
        self.workspace = Path(workspace)

    def run(self, path: str) -> str:  # type: ignore[override]
        target = resolve_in_workspace(self.workspace, path)
        if not target.exists():
            raise ToolError(f"{path!r} does not exist. Use list_files to see what is there.")
        if target.is_dir():
            raise ToolError(f"{path!r} is a folder - use list_files on it instead.")
        size = target.stat().st_size
        if size > MAX_READ_BYTES:
            raise ToolError(f"{path!r} is {size // 1024}KB, too large to read in full.")

        suffix = target.suffix.lower()
        if suffix == ".pdf":
            text = _read_pdf(target)
        elif suffix in {".docx", ".dotx"}:
            text = _read_docx(target)
        else:
            text = _read_text(target, path)

        if len(text) > MAX_CHARS:
            text = text[:MAX_CHARS] + f"\n\n[truncated at {MAX_CHARS} characters]"
        return f"{path}:\n\n{text}"


class ListFilesTool(Tool):
    name = "list_files"
    description = (
        "List the files in the student's homework folder so you can find the assignment. "
        "Returns paths relative to that folder, with sizes."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "subdirectory": {
                "type": "string",
                "description": "Optional subfolder to list. Defaults to the homework folder.",
            }
        },
    }

    def __init__(self, workspace: Path, max_entries: int = 200) -> None:
        self.workspace = Path(workspace)
        self.max_entries = max_entries

    def run(self, subdirectory: str = ".") -> str:  # type: ignore[override]
        root = resolve_in_workspace(self.workspace, subdirectory)
        if not root.exists():
            raise ToolError(f"{subdirectory!r} does not exist.")
        if not root.is_dir():
            raise ToolError(f"{subdirectory!r} is a file, not a folder.")

        entries: list[str] = []
        for path in sorted(root.rglob("*")):
            if any(part in SKIP_DIRS or part.startswith(".") for part in path.parts[len(root.parts):]):
                continue
            rel = path.relative_to(self.workspace.resolve())
            if path.is_dir():
                entries.append(f"{rel}/")
            else:
                entries.append(f"{rel}  ({path.stat().st_size} bytes)")
            if len(entries) >= self.max_entries:
                entries.append(f"[stopped at {self.max_entries} entries]")
                break
        if not entries:
            return f"{subdirectory} is empty."
        return "\n".join(entries)


class WriteFileTool(Tool):
    name = "write_file"
    description = (
        "Write a file in the student's homework folder - study notes, a summary of what you "
        "worked through, or code they asked you to save. Ask before overwriting something they "
        "wrote. Do not use this to hand in work on their behalf."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Path relative to the homework folder."},
            "content": {"type": "string", "description": "Full file contents."},
        },
        "required": ["path", "content"],
    }
    requires_approval = True

    def __init__(self, workspace: Path) -> None:
        self.workspace = Path(workspace)

    def run(self, path: str, content: str) -> str:  # type: ignore[override]
        target = resolve_in_workspace(self.workspace, path)
        if target.is_dir():
            raise ToolError(f"{path!r} is a folder.")
        existed = target.exists()
        target.parent.mkdir(parents=True, exist_ok=True)
        try:
            target.write_text(content, encoding="utf-8")
        except OSError as exc:
            raise ToolError(f"could not write {path!r}: {exc}") from None
        verb = "overwrote" if existed else "wrote"
        return f"{verb} {path} ({len(content)} characters)"
